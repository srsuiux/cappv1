"""One unreadable item must never stop a file: it is skipped, noted, and the rest finishes."""

import asyncio
import logging

from docx import Document
from pptx import Presentation

from cappv1 import events, outputs, pipeline
from cappv1.docx_reader import extract_pictures as docx_pictures
from cappv1.docx_reader import extract_tables, get_table_alt_text
from cappv1.pipeline import explain_error
from cappv1.services import files
from cappv1.services.runs import RunManager
from cappv1.state import load_state
from tests.helpers import alt_texts, make_deck
from tests.test_docx import docx_config, make_docx, run

BROKEN = "no relationship with key 'rId4'"


def break_first_picture(path):
    """A picture whose image part is gone, as in a damaged deck: python-pptx raises KeyError on it."""
    prs = Presentation(path)
    pic = next(s for s in prs.slides[0].shapes if s._element.tag.endswith("}pic"))
    prs.slides[0].part.rels.pop(pic._element.blip_rId)
    prs.save(path)


def collect():
    seen = []
    return seen, events.subscribe(lambda name, data: seen.append((name, data)))


def test_a_broken_picture_is_skipped_and_the_file_still_finishes(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    deck = make_deck(config.input_dir / "a.pptx")
    break_first_picture(deck)
    seen, unsubscribe = collect()
    assert asyncio.run(pipeline.run(config)) == []  # the file did not fail
    unsubscribe()

    state = load_state(config.state_dir, "a.pptx")
    assert state.status == "complete"
    assert [(s["where"], s["kind"]) for s in state.skipped] == [("slide 1", "picture")]
    assert state.skipped[0]["error"]
    assert (
        alt_texts(outputs.latest(config.output_dir, "a.pptx")).count(fake_ollama.reply) == 2
    )  # the table and chart on slide 2 were still done
    assert [d["where"] for n, d in seen if n == "element_skipped"] == ["slide 1"]
    saved = next(d for n, d in seen if n == "file_saved")
    assert saved["skipped"] == 1 and not saved["incomplete"]
    assert load_state(config.state_dir, "a.pptx").completed_slide_indices == [0, 1, 2]


def test_an_unreadable_slide_is_skipped_as_a_whole(tmp_path, fake_ollama, monkeypatch):
    def boom(slide):
        raise KeyError(BROKEN)

    monkeypatch.setattr(pipeline, "extract_tables", boom)
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    make_deck(config.input_dir / "a.pptx")
    assert asyncio.run(pipeline.run(config)) == []
    state = load_state(config.state_dir, "a.pptx")
    assert [s["where"] for s in state.skipped] == ["slide 1", "slide 2", "slide 3"]
    assert all(s["kind"] == "whole slide" for s in state.skipped)
    assert "broken link" in state.skipped[0]["error"]  # plain words, not a bare KeyError
    assert (outputs.latest(config.output_dir, "a.pptx")).exists() and state.status == "complete"


def test_a_broken_picture_in_a_word_document_is_skipped_too(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    path = make_docx(config.input_dir / "a.docx")
    doc = Document(str(path))
    rid = docx_pictures(doc)[0].doc_pr.getparent().xpath(".//a:blip/@r:embed")[0]
    doc.part.rels.pop(rid)
    doc.save(str(path))
    assert run(config) == []
    state = load_state(config.state_dir, "a.docx")
    assert state.status == "complete"
    assert [s["kind"] for s in state.skipped] == ["picture"] and "missing" in state.skipped[0]["error"]
    out = Document(str(outputs.latest(config.output_dir, "a.docx")))
    assert get_table_alt_text(extract_tables(out)[0].table) == fake_ollama.reply  # the table was still done


def test_notes_are_kept_for_the_files_list_and_start_fresh_each_run(tmp_path, configured, monkeypatch):
    deck = make_deck(tmp_path / "a.pptx")
    break_first_picture(deck)
    files.store_path(configured, deck)
    runs = RunManager(configured)
    runs.start()
    assert runs.wait(30)
    info = files.list_files(configured)[0]
    assert info.state == "done" and [s["kind"] for s in info.skipped] == ["picture"]
    assert info.as_dict()["skipped"][0]["where"] == "slide 1"
    runs.start(force=True)
    assert runs.wait(30)
    assert len(files.list_files(configured)[0].skipped) == 1  # rebuilt, not doubled


def test_a_browser_started_run_prints_its_log_in_the_terminal(configured, tmp_path):
    files.store_path(configured, make_deck(tmp_path / "a.pptx"))
    log = logging.getLogger("pptx_alt_text")
    runs = RunManager(configured)
    runs.start()
    assert runs.wait(30)
    assert not any(type(h) is logging.StreamHandler for h in log.handlers)  # the terminal session draws its own view
    runs.start(force=True, console_logs=True)
    assert runs.wait(30)
    assert any(type(h) is logging.StreamHandler for h in log.handlers)


def test_errors_are_explained_in_plain_words():
    assert "broken link" in explain_error(KeyError(BROKEN))
    assert explain_error(ValueError("bad value")) == "ValueError: bad value"
    assert explain_error(RuntimeError()) == "RuntimeError"


def test_a_null_model_reply_means_try_the_fallback_not_crash():
    import types

    import pytest

    from cappv1.llm_client import _ModelUnavailable, _parse_completion

    config = types.SimpleNamespace(provider="openrouter")
    with pytest.raises(_ModelUnavailable):
        _parse_completion(config, {"choices": [{"message": {"content": None}}]})
    assert _parse_completion(config, {"choices": [{"message": {"content": " hi "}}]}) == "hi"


def test_an_empty_model_reply_is_asked_again_before_giving_up(monkeypatch):
    import asyncio
    import types

    import httpx

    from cappv1 import llm_client

    replies = iter([None, "A diagram."])

    def handler(request):
        return httpx.Response(200, json={"choices": [{"message": {"content": next(replies)}}]})

    async def no_wait(_):
        return None

    monkeypatch.setattr(llm_client.asyncio, "sleep", no_wait)
    config = types.SimpleNamespace(
        provider="openrouter", base_url="http://x", api_key="k", request_timeout=5, max_alt_text_chars=125, ollama_base_url="", model="m"
    )
    monkeypatch.setattr(llm_client, "_build_request", lambda c, m, msgs: ("http://x/v1", {}, {}))

    async def go():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await llm_client._request_completion(client, config, "m", [], max_retries=3)

    assert asyncio.run(go()) == "A diagram."
