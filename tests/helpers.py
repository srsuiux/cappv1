from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def make_deck(path: Path, picture_color=(200, 60, 60)) -> Path:
    """A 3-slide deck: a picture, then a table and a chart, then an empty slide."""
    img = path.with_suffix(".png")
    Image.new("RGB", (300, 200), picture_color).save(img)
    prs = Presentation()
    for n in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"Slide {n + 1}"
        if n == 0:
            slide.shapes.add_picture(str(img), Inches(1), Inches(2), Inches(3))
        if n == 1:
            slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table.cell(0, 0).text = "a"
            data = CategoryChartData()
            data.categories = ["A", "B"]
            data.add_series("S", (1, 3))
            slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(4), Inches(2.5), data)
    prs.save(path)
    img.unlink()
    return path


def alt_texts(path: Path) -> list[str]:
    from cappv1.pptx_reader import get_alt_text

    return [t for slide in Presentation(path).slides for shape in slide.shapes if (t := get_alt_text(shape))]
