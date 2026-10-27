import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from src.audit import audit_deck, contrast_ratio
from src.services import audit as audit_service
from src.services.errors import ServiceError
from tests.helpers import make_deck


def _deck(path, text_rgb, fill_rgb=None, size=12):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    if fill_rgb:
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(fill_rgb)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Hello there"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(text_rgb)
    prs.save(path)
    return path


def test_contrast_ratio_matches_wcag():
    assert round(contrast_ratio("000000", "FFFFFF"), 1) == 21.0
    assert round(contrast_ratio("FFFFFF", "CA9910"), 1) == 2.6


def test_low_contrast_text_is_reported(tmp_path):
    report = audit_deck(_deck(tmp_path / "a.pptx", "FFFFFF", "CA9910"))
    assert [f.kind for f in report.findings] == ["contrast"] and "2.6:1" in report.findings[0].detail


def test_readable_text_passes(tmp_path):
    assert audit_deck(_deck(tmp_path / "a.pptx", "000000", "FFFFFF")).findings == []


def test_large_text_needs_less_contrast(tmp_path):
    grey = "949494"  # about 3.0:1 on white
    assert audit_deck(_deck(tmp_path / "small.pptx", grey, "FFFFFF", size=12)).findings
    assert not audit_deck(_deck(tmp_path / "large.pptx", grey, "FFFFFF", size=24)).findings


def test_text_over_a_template_shape_uses_that_shape(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(4), Inches(1))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string("CA9910")
    box = slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(2), Inches(0.5))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "On the band"
    run.font.color.rgb = RGBColor.from_string("FFFFFF")
    prs.save(tmp_path / "a.pptx")
    assert audit_deck(tmp_path / "a.pptx").findings[0].kind == "contrast"


def test_reading_order_flags_a_shape_read_before_one_above_it(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for text, top in (("lower", 4), ("upper", 1)):  # XML order: lower first
        slide.shapes.add_textbox(Inches(1), Inches(top), Inches(3), Inches(0.5)).text_frame.text = text
    prs.save(tmp_path / "a.pptx")
    findings = audit_deck(tmp_path / "a.pptx").findings
    assert [f.kind for f in findings] == ["reading_order"] and "lower" in findings[0].detail


def test_audit_never_changes_the_deck(tmp_path):
    path = _deck(tmp_path / "a.pptx", "FFFFFF", "CA9910")
    before = path.read_bytes()
    audit_deck(path)
    assert path.read_bytes() == before


def test_groups_collapse_the_same_cause(tmp_path):
    prs = Presentation()
    for _ in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string("CA9910")
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Same"
        run.font.color.rgb = RGBColor.from_string("FFFFFF")
    prs.save(tmp_path / "a.pptx")
    groups = audit_deck(tmp_path / "a.pptx").grouped()
    assert len(groups) == 1 and groups[0]["slides"] == [1, 2, 3]


def test_service_prefers_the_finished_copy_and_errors_plainly(ws, tmp_path):
    make_deck(ws.input_dir / "talk.pptx")
    assert audit_service.report(ws, "talk.pptx")["checked"] == "original"
    make_deck(ws.output_dir / "talk.pptx")
    assert audit_service.report(ws, "talk.pptx")["checked"] == "finished copy"
    with pytest.raises(ServiceError) as err:
        audit_service.report(ws, "missing.pptx")
    assert err.value.code == "not_found"


def test_table_cells_are_checked(tmp_path):
    from pptx.oxml.ns import qn

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    tbl = table._tbl
    for style in tbl.tblPr.findall(qn("a:tableStyleId")):
        tbl.tblPr.remove(style)  # an unstyled table: colors are only what the cells say
    cell = table.cell(0, 0)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string("CA9910")
    run = cell.text_frame.paragraphs[0].add_run()
    run.text = "Header"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("FFFFFF")
    prs.save(tmp_path / "a.pptx")
    report = audit_deck(tmp_path / "a.pptx")
    assert [f.kind for f in report.findings] == ["contrast"] and "Header" in report.findings[0].message


def test_text_in_a_group_uses_the_group_s_coordinates(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    band = slide.shapes.add_shape(1, Inches(5), Inches(5), Inches(3), Inches(1))  # dark band, far from the group's raw coords
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor.from_string("101010")
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(5.2), Inches(5.2), Inches(2), Inches(0.5))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "White on dark"
    run.font.color.rgb = RGBColor.from_string("FFFFFF")
    prs.save(tmp_path / "a.pptx")
    assert audit_deck(tmp_path / "a.pptx").findings == []
