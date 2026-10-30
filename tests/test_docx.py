from src import outputs
import asyncio
import io
import types
import zipfile

import pytest
from docx import Document
from docx.shared import Inches
from lxml import etree
from PIL import Image

from src import pipeline
from src.config import load_config
from src.docx_reader import (
    extract_pictures,
    extract_tables,
    get_alt_text,
    get_context,
    get_table_alt_text,
    mark_decorative,
    set_table_alt_text,
)

LONG = "A bar chart of quarterly sales for each region, used to compare the north and the south"

def png(color=(200, 60, 60), size=(300, 200)) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    buf.seek(0)
    return buf


def make_docx(path, *, tiny=False):
    doc = Document()
    doc.add_heading("Results", level=1)
    doc.add_paragraph("Sales rose in the north.")
    doc.add_picture(png(), width=Inches(2))
    doc.add_paragraph("Figure 1 shows the trend.")
    if tiny:
        doc.add_picture(png(size=(10, 10)), width=Inches(0.2))
    table = doc.add_table(rows=2, cols=2)
    for i, text in enumerate(["Region", "Sales", "North", "40"]):
        table.cell(i // 2, i % 2).text = text
    doc.save(path)
    return path


def docx_config(tmp_path, fake_ollama, monkeypatch):
    for k, v in {"PROVIDER": "ollama", "OLLAMA_BASE_URL": fake_ollama.url, "MODEL": "gemma3:4b"}.items():
        monkeypatch.setenv(k, v)
    config = load_config(tmp_path / "work")
    config.input_dir.mkdir(parents=True)
    return config


def run(config, **kw):
    return asyncio.run(pipeline.run(config, **kw))


def out_doc(config, name="a.docx"):
    return Document(str(outputs.latest(config.output_dir, name)))


def test_finds_pictures_and_tables_with_their_context(tmp_path):
    doc = Document(str(make_docx(tmp_path / "a.docx")))
    pics, tables = extract_pictures(doc), extract_tables(doc)
    assert len(pics) == 1 and len(tables) == 1
    ctx = get_context(doc, pics[0].block)
    assert ctx.title == "Results" and ctx.kind == "document"
    assert "Sales rose in the north." in ctx.body_text and "Figure 1" in ctx.body_text


def test_a_run_writes_alt_text_for_pictures_and_tables(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    make_docx(config.input_dir / "a.docx")
    assert run(config) == []
    doc = out_doc(config)
    assert get_alt_text(extract_pictures(doc)[0].doc_pr) == fake_ollama.reply
    assert get_table_alt_text(extract_tables(doc)[0].table) == fake_ollama.reply
    assert fake_ollama.calls == 2
    # finished and unchanged: a second run does nothing
    run(config)
    assert fake_ollama.calls == 2


def test_only_alt_text_changes_in_the_document(tmp_path, fake_ollama, monkeypatch):
    """The skills.md invariant: everything but the alt-text fields comes out the same."""
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    src = make_docx(config.input_dir / "a.docx")
    run(config)

    def canon(zf, name):
        root = etree.fromstring(zf.read(name))
        for el in root.iter():
            el.attrib.pop("descr", None)
        for tag in ("tblDescription", "tblHeader", "tblLook"):  # the fields the tool is allowed to write
            for el in list(root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + tag)):
                el.getparent().remove(el)
        for el in list(root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}trPr")):
            if len(el) == 0:
                el.getparent().remove(el)
        return etree.tostring(root, method="c14n")

    with zipfile.ZipFile(src) as a, zipfile.ZipFile(outputs.latest(config.output_dir, "a.docx")) as b:
        assert sorted(a.namelist()) == sorted(b.namelist())
        for name in a.namelist():
            if name == "docProps/core.xml":
                continue  # the Title is filled in when it is empty or names something else
            if name.endswith((".xml", ".rels")):
                assert canon(a, name) == canon(b, name), name
            else:
                assert a.read(name) == b.read(name), name


def test_the_package_keeps_the_originals_order_and_bytes(tmp_path, fake_ollama, monkeypatch):
    """Only word/document.xml is rewritten; Blackboard's previewer refuses repackaged files."""
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    src = make_docx(config.input_dir / "a.docx")
    run(config)

    with zipfile.ZipFile(src) as a, zipfile.ZipFile(outputs.latest(config.output_dir, "a.docx")) as b:
        assert a.namelist() == b.namelist()
        for name in a.namelist():
            if name not in ("word/document.xml", "docProps/core.xml"):  # the Title is filled in when it is wrong
                assert a.read(name) == b.read(name), name
        assert b.read("word/document.xml").startswith(b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>')


def set_mode(path, mode):
    """Rewrites the compatibility mode in a Word file's settings, the way an old save has it."""
    with zipfile.ZipFile(path) as z:
        parts = {i: z.read(i) for i in z.namelist()}
    xml = parts["word/settings.xml"].decode()
    if "compatibilityMode" in xml:
        import re
        xml = re.sub(r'(compatibilityMode"[^>]*w:val=")\d+', rf"\g<1>{mode}", xml)
    else:
        xml = xml.replace("</w:settings>", f'<w:compat><w:compatSetting w:name="compatibilityMode" '
                          f'w:uri="http://schemas.microsoft.com/office/word" w:val="{mode}"/></w:compat></w:settings>')
    parts["word/settings.xml"] = xml.encode()
    with zipfile.ZipFile(path, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)


def test_tables_in_an_old_compatibility_mode_document_are_left_alone(tmp_path, fake_ollama, monkeypatch):
    """Word cannot open a mode 11 file that has table alt text, so pictures only, and a note."""
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    set_mode(make_docx(config.input_dir / "a.docx"), 11)
    run(config)
    xml = zipfile.ZipFile(outputs.latest(config.output_dir, "a.docx")).read("word/document.xml").decode()
    assert "tblDescription" not in xml and "tblCaption" not in xml
    assert 'descr="' in xml and 'descr=""' not in xml  # the picture still got its text
    from src.state import load_state
    notes = load_state(config.state_dir, "a.docx").skipped
    assert len(notes) == 1 and notes[0]["kind"] == "tables"
    assert notes[0]["error"].startswith("1 table not described") and "compatibility mode (11)" in notes[0]["error"]
    assert "Error" not in notes[0]["error"]


def test_tables_in_a_modern_document_are_described(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    set_mode(make_docx(config.input_dir / "a.docx"), 15)
    run(config)
    assert "tblDescription" in zipfile.ZipFile(outputs.latest(config.output_dir, "a.docx")).read("word/document.xml").decode()


def add_object(path, *, alt=None):
    """Embeds a small Word file as an OLE object, the way a Visio or Excel one sits in a document."""
    inner = io.BytesIO()
    inside = Document()
    inside.add_paragraph("Customer places Order; Order contains Product")
    inside.save(inner)
    alt_attr = f' alt="{alt}"' if alt else ""
    ns = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
          'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
    obj = (f'<w:p {ns}><w:r><w:object w:dxaOrig="100" w:dyaOrig="100"><v:shape id="ole_rId9" style="width:10pt"{alt_attr}>'
           f'<v:imagedata o:title=""/></v:shape><o:OLEObject Type="Embed" ProgID="" ShapeID="ole_rId9" '
           f'DrawAspect="Content" ObjectID="_1" r:id="rId9"/></w:object></w:r></w:p>')
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts["word/embeddings/oleObject1.bin"] = inner.getvalue()
    rels = parts["word/_rels/document.xml.rels"].decode().replace(
        "</Relationships>", '<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        'relationships/oleObject" Target="embeddings/oleObject1.bin"/></Relationships>')
    parts["word/_rels/document.xml.rels"] = rels.encode()
    types = parts["[Content_Types].xml"].decode().replace(
        "<Override", '<Default Extension="bin" ContentType="application/vnd.openxmlformats-officedocument.oleObject"/><Override', 1)
    parts["[Content_Types].xml"] = types.encode()
    doc = parts["word/document.xml"].decode().replace("<w:sectPr", obj.replace(f" {ns}", "") + "<w:sectPr", 1)
    parts["word/document.xml"] = doc.encode()
    with zipfile.ZipFile(path, "w") as z:
        for name, data in parts.items():
            z.writestr(name, data)


def object_alts(config, name="a.docx"):
    xml = etree.fromstring(zipfile.ZipFile(outputs.latest(config.output_dir, name)).read("word/document.xml"))
    return [v.get("alt") for v in xml.iter("{urn:schemas-microsoft-com:vml}shape")]


def test_an_embedded_object_gets_alt_text_on_its_preview_shape(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    src = make_docx(config.input_dir / "a.docx")
    add_object(src)
    run(config)
    assert [a for a in object_alts(config) if a]  # described, not left bare
    assert "ProgID" in zipfile.ZipFile(outputs.latest(config.output_dir, "a.docx")).read("word/document.xml").decode()


def test_an_authors_object_alt_text_is_kept(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    src = make_docx(config.input_dir / "a.docx")
    add_object(src, alt="ER diagram drawn by the author")
    fake_ollama.reply = "KEEP"
    run(config)
    assert object_alts(config) == ["ER diagram drawn by the author"]


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def listing_free_doc(tmp_path, rows):
    """A document holding one table made of `rows`, the way LibreOffice writes it: no style, no look."""
    doc = Document()
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    pr = table._tbl.tblPr
    for tag in ("tblStyle", "tblLook"):
        for el in pr.findall(W + tag):
            pr.remove(el)
    return table, pr


def test_a_plain_table_gets_its_first_row_marked_as_the_header(tmp_path):
    from src.docx_reader import mark_table_header

    table, pr = listing_free_doc(tmp_path, [["Tool", "Vendor"], ["Erwin", "Quest"]])
    assert mark_table_header(table) == "set"
    assert pr.find(W + "tblLook").get(W + "firstRow") == "1"
    first = table._tbl.findall(W + "tr")[0]
    assert first.find(f"{W}trPr/{W}tblHeader") is not None
    assert table._tbl.findall(W + "tr")[1].find(f"{W}trPr") is None
    assert mark_table_header(table) == "already"


def test_header_look_goes_before_an_existing_description(tmp_path):
    from src.docx_reader import mark_table_header

    table, pr = listing_free_doc(tmp_path, [["A", "B"], ["1", "2"]])
    etree.SubElement(pr, W + "tblDescription").set(W + "val", "x")
    mark_table_header(table)
    tags = [etree.QName(c).localname for c in pr]
    assert tags.index("tblLook") < tags.index("tblDescription")


@pytest.mark.parametrize("rows", [
    [["Only", "row"]],
    [["A"], ["1"]],
    [["A", ""], ["1", "2"]],
])
def test_every_table_gets_a_header_row_whatever_it_holds(tmp_path, rows):
    """Ally flags a table without a header row, including one-row boxes and one-column listings."""
    from src.docx_reader import mark_table_header

    table, pr = listing_free_doc(tmp_path, rows)
    assert mark_table_header(table) == "set"
    assert table._tbl.findall(W + "tr")[0].find(f"{W}trPr/{W}tblHeader") is not None


def test_a_styled_table_gets_the_repeat_setting_but_not_the_look_box(tmp_path):
    from src.docx_reader import mark_table_header

    table, pr = listing_free_doc(tmp_path, [["A", "B"], ["1", "2"]])
    etree.SubElement(pr, W + "tblStyle").set(W + "val", "Grid")
    assert mark_table_header(table) == "set"
    assert pr.find(W + "tblLook") is None  # the box could switch on the style's first-row formatting
    assert table._tbl.findall(W + "tr")[0].find(f"{W}trPr/{W}tblHeader") is not None


def test_header_rows_are_marked_even_when_table_descriptions_are_skipped(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    src = make_docx(config.input_dir / "a.docx")
    set_mode(src, 11)
    run(config)
    xml = zipfile.ZipFile(outputs.latest(config.output_dir, "a.docx")).read("word/document.xml").decode()
    assert "tblDescription" not in xml and "<w:tblHeader" in xml


def test_the_decorative_flag_uses_the_standard_prefix(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    make_docx(config.input_dir / "a.docx", tiny=True)
    run(config)
    with zipfile.ZipFile(outputs.latest(config.output_dir, "a.docx")) as z:
        xml = z.read("word/document.xml").decode()
    assert "<a:extLst" in xml and "ns0:" not in xml


def test_tiny_pictures_are_marked_decorative_without_a_model_call(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    make_docx(config.input_dir / "a.docx", tiny=True)
    run(config)
    pics = extract_pictures(out_doc(config))
    tiny = next(p for p in pics if len(p.image_bytes) < 500)
    assert get_alt_text(tiny.doc_pr) == ""
    assert tiny.doc_pr.xpath(".//*[local-name()='decorative']")
    assert fake_ollama.calls == 2  # the big picture and the table only


def test_the_authors_alt_text_is_kept_when_the_model_says_keep(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    src = make_docx(config.input_dir / "a.docx")
    doc = Document(str(src))
    pic = extract_pictures(doc)[0]
    pic.doc_pr.set("descr", LONG)
    set_table_alt_text(extract_tables(doc)[0].table, LONG)
    doc.save(str(src))
    fake_ollama.reply = "KEEP"
    run(config)
    out = out_doc(config)
    assert get_alt_text(extract_pictures(out)[0].doc_pr) == LONG
    assert get_table_alt_text(extract_tables(out)[0].table) == LONG


def test_table_description_goes_in_schema_order(tmp_path):
    doc = Document(str(make_docx(tmp_path / "a.docx")))
    table = extract_tables(doc)[0].table
    pr = table._tbl.tblPr
    etree.SubElement(pr, "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblPrChange")
    set_table_alt_text(table, "x")
    tags = [etree.QName(c).localname for c in pr]
    assert tags.index("tblDescription") < tags.index("tblPrChange")


def test_decorative_marking_is_not_added_twice(tmp_path):
    doc = Document(str(make_docx(tmp_path / "a.docx")))
    doc_pr = extract_pictures(doc)[0].doc_pr
    mark_decorative(doc_pr)
    mark_decorative(doc_pr)
    assert len(doc_pr.xpath(".//*[local-name()='decorative']")) == 1


def test_a_failed_model_call_leaves_the_file_to_retry(tmp_path, fake_ollama, monkeypatch):
    async def no_wait(_):
        pass

    monkeypatch.setattr("src.llm_client.asyncio.sleep", no_wait)
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    make_docx(config.input_dir / "a.docx")
    fake_ollama.stop()
    fake_ollama.server.server_close()  # every call now fails to connect
    run(config)
    from src.state import load_state

    state = load_state(config.state_dir, "a.docx")
    assert state.status == "in_progress" and state.completed_slide_indices == []
    assert (outputs.latest(config.output_dir, "a.docx")).exists()  # saved as far as it got


# ---- services, runs and the web server --------------------------------------------
def test_docx_uploads_are_accepted_and_checked(ws, tmp_path):
    from src.services import files
    from src.services.errors import ServiceError

    info = files.store_path(ws, make_docx(tmp_path / "paper.docx"))
    assert info.name == "paper.docx" and info.state == "new"
    assert [f.name for f in files.list_files(ws)] == ["paper.docx"]
    with pytest.raises(ServiceError) as err:
        files.store(ws, "fake.docx", [b"not a zip"])
    assert err.value.code == "unsupported_file"
    # a deck renamed to .docx is not a Word document
    from tests.helpers import make_deck

    with pytest.raises(ServiceError):
        files.store_path(ws, make_deck(tmp_path / "d.pptx").rename(tmp_path / "d.docx"))
    assert files.safe_name("Paper.DOCX") == "Paper.docx"


def test_word_lock_files_are_ignored(ws, tmp_path):
    from src.services import files

    files.store_path(ws, make_docx(tmp_path / "paper.docx"))
    (ws.input_dir / "~$paper.docx").write_bytes(b"lock")
    assert [f.name for f in files.list_files(ws)] == ["paper.docx"]


def test_a_run_through_the_service_captions_a_docx(configured, tmp_path):
    from src.services import files
    from src.services.runs import RunManager

    files.store_path(configured, make_docx(tmp_path / "paper.docx"))
    runs = RunManager(configured)
    runs.start()
    assert runs.wait(30)
    assert runs.status().state == "done" and runs.status().failed_files == []
    assert files.list_files(configured)[0].state == "done"
    assert files.output_file(configured, "paper.docx").is_file()
    assert get_alt_text(extract_pictures(out_doc_at(outputs.latest(configured.output_dir, "paper.docx")))[0].doc_pr)


def out_doc_at(path):
    return Document(str(path))


def test_the_accessibility_check_says_it_is_decks_only(configured, tmp_path):
    from src.services import audit, files
    from src.services.errors import ServiceError

    files.store_path(configured, make_docx(tmp_path / "paper.docx"))
    with pytest.raises(ServiceError) as err:
        audit.report(configured, "paper.docx")
    assert err.value.code == "audit_unsupported"


def test_clean_removes_documents_too(configured, tmp_path):
    from src.services import files

    files.store_path(configured, make_docx(tmp_path / "paper.docx"))
    assert files.clean(configured)["input"] == 1
    assert files.list_files(configured) == []


def test_run_events_say_what_kind_of_file_it_is(configured, tmp_path):
    """The web page shows slide tiles for a deck and item chips for a Word document."""
    from src.services import files
    from src.services.runs import RunManager
    from tests.helpers import make_deck

    files.store_path(configured, make_docx(tmp_path / "paper.docx"))
    files.store_path(configured, make_deck(tmp_path / "talk.pptx"))
    runs = RunManager(configured)
    runs.start()
    assert runs.wait(30)
    events = runs.events_since(0)
    kinds = {e["filename"]: e["kind"] for e in events if e["event"] == "file_start"}
    assert kinds == {"paper.docx": "document", "talk.pptx": "slides"}
    opened = {e["filename"]: e for e in events if e["event"] == "file_open"}
    assert opened["paper.docx"]["kind"] == "document" and opened["paper.docx"]["items"] == 2
    assert opened["talk.pptx"]["kind"] == "slides" and opened["talk.pptx"]["total_slides"] == 3
    written = [e for e in events if e["event"] == "alt_written" and e["filename"] == "paper.docx"]
    assert len(written) == 2


def test_an_embedded_excel_sheet_is_read_not_mistaken_for_a_word_file():
    from src.pptx_reader import ole_object_to_text

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/sharedStrings.xml", '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   "<si><t>Quarter</t></si><si><r><t>Revenue</t></r></si></sst>")
    text = ole_object_to_text(buf.getvalue())
    assert "Quarter" in text and "Revenue" in text
    assert ole_object_to_text(b"PK\x03\x04 not really a zip") == ""


def test_the_terminal_says_document_not_slide_for_a_word_file():
    import io as _io

    from rich.console import Console

    from src.tui.runview import RunView

    out = _io.StringIO()
    view = RunView(Console(file=out, width=100, force_terminal=False), types.SimpleNamespace())
    view.on_event("file_start", {"filename": "a.docx", "index": 1, "total": 1})
    view.on_event("file_open", {"total_slides": 1, "resumed": 0, "kind": "document"})
    view.on_event("slide_end", {"slide": 0, "empty": False, "captioned": 4, "enhanced": 0, "kept": 0,
                                "decorative": 0, "cached": 0, "failed": 0})
    assert "document" in out.getvalue() and "slide" not in out.getvalue()


# ---- structure fixes Ally scores -------------------------------------------------------------

def styled(doc, text, style):
    return doc.add_paragraph(text, style=style)


def bare(doc):
    """Drops the template styles built on Heading 1-9, as in the converted documents, which have none."""
    for st in list(doc.styles.element.iter(W + "style")):
        based = st.find(W + "basedOn")
        if based is not None and based.get(W + "val").startswith("Heading") and not st.get(W + "styleId").startswith("Heading"):
            st.getparent().remove(st)
    return doc


def rpr_of(styles_el, style_id):
    st = next(s for s in styles_el.iter(W + "style") if s.get(W + "styleId") == style_id)
    rpr = st.find(W + "rPr")
    return etree.tostring(rpr, method="c14n") if rpr is not None else b""


def test_heading_levels_are_renumbered_without_changing_how_they_look():
    from src.docx_structure import normalize_headings

    doc = bare(Document())
    before = rpr_of(doc.styles.element, "Heading4"), rpr_of(doc.styles.element, "Heading7")
    styled(doc, "Chapter", "Heading 4")
    styled(doc, "Section", "Heading 7")
    styled(doc, "Another section", "Heading 7")
    assert normalize_headings(doc)
    ids = [p.style.style_id for p in doc.paragraphs]
    assert ids == ["Heading1", "Heading2", "Heading2"]
    assert (rpr_of(doc.styles.element, "Heading1"), rpr_of(doc.styles.element, "Heading2")) == before
    outline = next(s for s in doc.styles.element.iter(W + "style") if s.get(W + "styleId") == "Heading2")
    assert outline.find(f"{W}pPr/{W}outlineLvl").get(W + "val") == "1"


def test_a_document_that_already_starts_at_heading_1_is_not_touched():
    from src.docx_structure import normalize_headings

    doc = Document()
    styled(doc, "One", "Heading 1")
    styled(doc, "Two", "Heading 2")
    assert normalize_headings(doc) == ""


def test_an_empty_heading_keeps_its_look_but_stops_being_a_heading():
    from src.docx_structure import normalize_headings

    doc = bare(Document())
    styled(doc, "One", "Heading 1")
    styled(doc, "", "Heading 2")
    assert "empty heading" in normalize_headings(doc)
    plain = doc.paragraphs[1].style
    assert plain.style_id == "Heading2Plain" and not plain.name.lower().startswith("heading")
    assert rpr_of(doc.styles.element, "Heading2Plain") == rpr_of(doc.styles.element, "Heading2")


def test_captions_become_heading_1_in_a_document_with_no_headings():
    from src.docx_structure import normalize_headings

    doc = Document()
    styled(doc, "Table 2-1: Quality", "Normal")
    styled(doc, "body text", "Normal")
    styled(doc, "Figure 2.3: Phases", "Normal")
    assert "captions" in normalize_headings(doc)
    assert [p.style.style_id for p in doc.paragraphs] == ["Heading1", "Normal", "Heading1"]
    h1 = next(s for s in doc.styles.element.iter(W + "style") if s.get(W + "styleId") == "Heading1")
    assert h1.find(W + "basedOn").get(W + "val") == "Normal" and h1.find(f"{W}pPr/{W}outlineLvl") is not None


def test_low_contrast_text_is_darkened_to_pass_and_keeps_its_hue():
    from docx.oxml.ns import qn

    from src.docx_check import contrast
    from src.docx_structure import fix_contrast

    doc = Document()
    run = doc.add_paragraph().add_run("light green")
    run._r.get_or_add_rPr()
    color = etree.SubElement(run._r.rPr, qn("w:color"))
    color.set(qn("w:val"), "70AD47")
    ok = doc.add_paragraph().add_run("black")
    etree.SubElement(ok._r.get_or_add_rPr(), qn("w:color")).set(qn("w:val"), "000000")
    notes = fix_contrast(doc)
    assert len(notes) == 1 and "70AD47" in notes[0]
    now = color.get(qn("w:val"))
    assert contrast(now, "FFFFFF") >= 7
    assert int(now[2:4], 16) > int(now[0:2], 16)  # still green
    assert fix_contrast(doc) == []


def test_the_title_is_filled_when_empty_or_wrong_and_left_alone_when_right():
    from src.docx_structure import set_title

    doc = Document()
    doc.add_paragraph("Chapter 2 Figures and Tables")
    doc.core_properties.title = "Table 14-4: from another chapter"
    assert set_title(doc) == "Chapter 2 Figures and Tables"
    assert doc.core_properties.title == "Chapter 2 Figures and Tables"
    assert set_title(doc) is None


def test_the_finished_file_passes_the_ally_self_check(tmp_path, fake_ollama, monkeypatch):
    from src.docx_check import check

    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    make_docx(config.input_dir / "a.docx")
    run(config)
    blocking = [i for i in check(outputs.latest(config.output_dir, "a.docx")) if i.severity in ("severe", "major")]
    assert blocking == [], [str(i) for i in blocking]


def table_with_heading_cells(doc, style="Heading 4"):
    table = doc.add_table(rows=2, cols=2)
    for c, text in enumerate(["Data Type", "Description"]):
        table.cell(0, c).paragraphs[0].text = text
        table.cell(0, c).paragraphs[0].style = doc.styles[style]
    table.cell(1, 0).text, table.cell(1, 1).text = "CHAR", "text"
    return table


def test_heading_styled_table_labels_do_not_count_as_headings_so_captions_are_promoted():
    """Ally ignores headings inside table cells; in four real chapters that was the only 'heading'."""
    from src.docx_structure import normalize_headings

    doc = bare(Document())
    styled(doc, "Table 3-1: Data types", "Normal")
    table_with_heading_cells(doc)
    styled(doc, "Figure 3.2: Layout", "Normal")
    assert "captions" in normalize_headings(doc)
    body_styles = [p.style.style_id for p in doc.paragraphs]
    assert body_styles == ["Heading1", "Heading1"]
    cells = [c.paragraphs[0].style.style_id for c in doc.tables[0].rows[0].cells]
    assert cells == ["Heading4", "Heading4"]  # the labels keep their style and look


def test_body_headings_are_renumbered_while_table_labels_keep_their_old_style():
    from src.docx_structure import normalize_headings

    doc = bare(Document())
    styled(doc, "Intro", "Heading 4")
    table_with_heading_cells(doc, "Heading 4")
    assert "renumbered" in normalize_headings(doc)
    assert doc.paragraphs[0].style.style_id == "Heading1"
    assert [c.paragraphs[0].style.style_id for c in doc.tables[0].rows[0].cells] == ["Heading4", "Heading4"]
    assert any(s.get(W + "styleId") == "Heading4" for s in doc.styles.element.iter(W + "style"))


def test_the_self_check_reports_headings_that_exist_only_inside_tables(tmp_path):
    from src.docx_check import check

    doc = Document()
    for i in range(14):
        doc.add_paragraph(f"body paragraph {i}")
    table_with_heading_cells(doc)
    doc.save(tmp_path / "t.docx")
    headings = [i for i in check(tmp_path / "t.docx") if i.item == "headings"]
    assert len(headings) == 1 and "inside table cells" in headings[0].message


def test_a_picture_no_model_can_read_is_described_by_its_caption(tmp_path, fake_ollama, monkeypatch):
    config = docx_config(tmp_path, fake_ollama, monkeypatch)
    src = make_docx(config.input_dir / "a.docx")
    with zipfile.ZipFile(src) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    media = next(n for n in parts if n.startswith("word/media/"))
    renamed = media.rsplit(".", 1)[0] + ".wmf"
    parts[renamed] = parts.pop(media)
    parts["word/_rels/document.xml.rels"] = parts["word/_rels/document.xml.rels"].replace(
        media.encode()[len("word/"):], renamed.encode()[len("word/"):])
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"<Default", b'<Default Extension="wmf" ContentType="image/x-wmf"/><Default', 1)
    doc_xml = parts["word/document.xml"].decode().replace("Figure 1 shows the trend.", "Figure 7.3: Hospital layout")
    caption = '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:r><w:t>Figure 7.3: Hospital layout</w:t></w:r></w:p>'
    parts["word/document.xml"] = doc_xml.encode()
    with zipfile.ZipFile(src, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    run(config)
    root = etree.fromstring(zipfile.ZipFile(outputs.latest(config.output_dir, "a.docx")).read("word/document.xml"))
    descrs = [d.get("descr") for d in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")]
    assert descrs == ["Figure 7.3: Hospital layout"], descrs


def test_what_counts_as_a_caption():
    from src.docx_reader import CAPTION

    for yes in ("Figure 7.3: Hospital layout", "Table 2-1. Quality", "Figure 2.9 ER Modeler Window", "Fig. 4 - Plan"):
        assert CAPTION.match(yes), yes
    for no in ("Figure 1 shows the trend.", "Table of contents", "The figure 3 above", "Table 5 is wide"):
        assert not CAPTION.match(no), no
