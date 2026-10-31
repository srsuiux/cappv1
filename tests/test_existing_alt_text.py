from pptx import Presentation

from cappv1 import outputs
from cappv1.pptx_reader import get_alt_text, set_alt_text
from cappv1.services import files
from cappv1.services.runs import RunManager
from tests.helpers import make_deck

LONG = "A chart of quarterly sales for each region, used to compare the north and the south"


def deck_with_picture_alt(ws, tmp_path, alt):
    path = make_deck(tmp_path / "a.pptx")
    prs = Presentation(path)
    for shape in prs.slides[0].shapes:
        if shape.shape_type == 13:
            set_alt_text(shape, alt)
    prs.save(path)
    files.store_path(ws, path)


def run_and_get_picture_alt(ws):
    runs = RunManager(ws)
    runs.start()
    assert runs.wait(30)
    slide_one = next(e for e in runs.events_since(0) if e["event"] == "slide_end" and e["slide"] == 0)
    pic = next(s for s in Presentation(outputs.latest(ws.output_dir, "a.pptx")).slides[0].shapes if s.shape_type == 13)
    return get_alt_text(pic), slide_one


def test_a_file_name_is_not_treated_as_the_authors_alt_text(configured, tmp_path):
    deck_with_picture_alt(configured, tmp_path, "image1.png")
    _, counts = run_and_get_picture_alt(configured)
    assert counts["captioned"] == 1 and counts["enhanced"] == 0 and counts["kept"] == 0


def test_short_author_text_is_improved_not_ignored(configured, fake_ollama, tmp_path):
    fake_ollama.reply = "Logo - the company logo in the header."
    deck_with_picture_alt(configured, tmp_path, "Logo")
    alt, counts = run_and_get_picture_alt(configured)
    assert counts["enhanced"] == 1 and alt == "Logo - the company logo in the header."


def test_adequate_author_text_is_kept_exactly_as_written(configured, fake_ollama, tmp_path):
    fake_ollama.reply = "KEEP"
    deck_with_picture_alt(configured, tmp_path, LONG)
    alt, counts = run_and_get_picture_alt(configured)
    assert counts["kept"] == 1 and alt == LONG


def test_force_replaces_the_authors_text(configured, fake_ollama, tmp_path):
    fake_ollama.reply = "A fresh description."
    deck_with_picture_alt(configured, tmp_path, LONG)
    runs = RunManager(configured)
    runs.start(force=True)
    assert runs.wait(30)
    pic = next(s for s in Presentation(outputs.latest(configured.output_dir, "a.pptx")).slides[0].shapes if s.shape_type == 13)
    assert get_alt_text(pic) == "A fresh description."


def test_different_images_never_share_a_caption(configured, fake_ollama, tmp_path):
    """Two flat-colored pictures look 'similar' to a perceptual hash; they must not collide."""
    for color, name in (((200, 60, 60), "red.pptx"), ((10, 200, 10), "green.pptx")):
        files.store_path(configured, make_deck(tmp_path / name, color))
    runs = RunManager(configured)
    runs.start()
    assert runs.wait(30)
    pictures = sum(1 for e in runs.events_since(0) if e["event"] == "element_start" and e["kind"] == "picture")
    assert pictures == 2  # both pictures hit the model; neither reused the other's caption


# ---- table header row ---------------------------------------------------------
def _table(rows, style=True, merge=False, blank=False):
    from pptx import Presentation
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape = slide.shapes.add_table(rows, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    pr = shape.table._tbl.tblPr
    for a in ("firstRow", "bandRow"):
        pr.attrib.pop(a, None)
    if not style:
        for child in pr.findall(qn("a:tableStyleId")):
            pr.remove(child)
    for i, c in enumerate(shape.table.rows[0].cells):
        c.text = "" if blank and i == 1 else f"h{i}"
    if merge:
        shape.table.cell(0, 0).merge(shape.table.cell(0, 1))
    return shape


def test_header_row_is_flagged_when_it_changes_nothing_visually():
    from cappv1.pptx_reader import mark_table_header

    shape = _table(3, style=False)
    assert mark_table_header(shape) == "set"
    assert shape.table._tbl.tblPr.get("firstRow") == "1"
    assert mark_table_header(shape) == "already"


def test_header_row_is_left_alone_when_unsafe():
    from cappv1.pptx_reader import mark_table_header

    for shape in (_table(3, style=True), _table(1, style=False), _table(3, style=False, merge=True), _table(3, style=False, blank=True)):
        assert mark_table_header(shape).startswith("skipped")
        assert shape.table._tbl.tblPr.get("firstRow") is None


def test_the_deck_keeps_every_part_except_the_slides_it_wrote_to(configured, tmp_path):
    """python-pptx drops a relationship whose target it cannot resolve (an embedded object with
    Target="NULL"), leaving the slide pointing at nothing; PowerPoint then asks to repair the file."""
    import zipfile

    src = make_deck(tmp_path / "a.pptx")
    broken = tmp_path / "broken.pptx"
    rel = ('<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
           'relationships/oleObject" Target="NULL"/>')
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(broken, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "ppt/slides/_rels/slide1.xml.rels":
                data = data.replace(b"</Relationships>", rel.encode() + b"</Relationships>")
            zout.writestr(info.filename, data)
    broken.replace(src)
    files.store_path(configured, src)
    runs = RunManager(configured)
    runs.start()
    assert runs.wait(30)
    with zipfile.ZipFile(src) as before, zipfile.ZipFile(outputs.latest(configured.output_dir, "a.pptx")) as after:
        assert after.namelist() == before.namelist()
        assert b'Id="rId99"' in after.read("ppt/slides/_rels/slide1.xml.rels")
        changed = [n for n in before.namelist() if before.read(n) != after.read(n)]
        assert changed and all(n.startswith("ppt/slides/slide") and n.endswith(".xml") for n in changed)
